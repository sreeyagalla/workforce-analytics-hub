"""Automated reporting: an Excel scorecard workbook and a PowerPoint deck.

Both are built from `collect()`, which computes everything through the governed
metric layer (so suppression and definitions are identical to the dashboard).
Narrative sentences are templates filled with computed values via metrics.fmt.
"""
from __future__ import annotations

import io
from dataclasses import dataclass
from datetime import date
from pathlib import Path

import numpy as np
import pandas as pd

from . import charts, forecast, metrics, scorecard, settings
from .metrics import MetricQuery, fmt, fmt_delta

SUPPRESSED_TEXT = "n<{k}"
METHOD_LABEL = {"last_rate": "end-of-year headcount x last year's turnover rate",
                "trailing_mean": "end-of-year headcount x mean turnover rate of up to 3 prior years",
                "naive_count": "same number of separations as last year"}


@dataclass
class ReportData:
    fy: int
    is_sample: bool
    n_agencies: int
    agency_card: scorecard.Scorecard
    location_card: scorecard.Scorecard
    trend: pd.DataFrame
    benchmark: pd.DataFrame
    watch: list[dict]
    outlook: pd.DataFrame
    backtest: pd.DataFrame
    driver: dict | None
    dq: pd.DataFrame
    headline: dict


def _org_trend(con) -> pd.DataFrame:
    out = None
    for key in settings.metric_catalog():
        df = metrics.compute(con, MetricQuery(key)).df[["fiscal_year", "value"]].rename(columns={"value": key})
        out = df if out is None else out.merge(df, on="fiscal_year", how="outer")
    return out.sort_values("fiscal_year").reset_index(drop=True)


def _benchmark(con) -> pd.DataFrame:
    tables = {r[0] for r in con.execute("SHOW TABLES").fetchall()}
    if "bls_benchmark" not in tables:
        return pd.DataFrame(columns=["fiscal_year", "value"])
    return con.execute("""SELECT fiscal_year, rate AS value FROM bls_benchmark
                          WHERE measure = 'total_separations_rate' ORDER BY 1""").df()


def _driver(con, unit: str, fy: int) -> dict | None:
    """For the unit with the highest turnover: turnover by pay basis (descriptive, not causal)."""
    code = next((c for c, v in settings.sources()["agencies"].items() if v["name"] == unit), None)
    if code is None:
        return None
    df = metrics.compute(con, MetricQuery("turnover_rate", group_by=["pay_basis"], filters={"agency": [code]},
                                          fiscal_years=[fy])).df
    df = df[~df["suppressed"] & df["value"].notna()].nlargest(2, "population")  # the two largest groups
    if len(df) < 2:
        return None
    df = df.sort_values("value", ascending=False)
    return {"unit": unit, "hi": df.iloc[0].to_dict(), "lo": df.iloc[1].to_dict()}


def collect(con, fy: int, is_sample: bool = False) -> ReportData:
    agency = scorecard.build(con, fy, "agency")
    location = scorecard.build(con, fy, "location")
    trend = _org_trend(con)
    watch = scorecard.watch_list(agency)
    red_turnover = [w for w in watch if w["metric"] == "turnover_rate"]
    driver = _driver(con, red_turnover[0]["unit"], fy) if red_turnover else None
    preds, summary = forecast.backtest(con)
    prev = trend[trend.fiscal_year == fy - 1]
    cur = trend[trend.fiscal_year == fy]
    headline = {k: (cur[k].iloc[0] if len(cur) else np.nan, prev[k].iloc[0] if len(prev) else np.nan)
                for k in settings.metric_catalog()}
    return ReportData(fy=fy, is_sample=is_sample, n_agencies=len(settings.sources()["agencies"]),
                      agency_card=agency, location_card=location, trend=trend, benchmark=_benchmark(con),
                      watch=watch, outlook=forecast.outlook(con), backtest=summary, driver=driver,
                      dq=con.execute("SELECT * FROM dq_results").df(), headline=headline)


# ------------------------------------------------------------------ narrative (templated, grounded)

def narrative(d: ReportData) -> dict[str, str]:
    cat = settings.metric_catalog()
    hc, hc_prev = d.headline["headcount"]
    t, t_prev = d.headline["turnover_rate"]
    n_red = len({w["unit"] for w in d.watch})
    out = {
        "headline": (f"End-of-year headcount was {fmt(hc, 'count')} in FY{d.fy} "
                     f"({(hc / hc_prev - 1) * 100:+.1f}% vs FY{d.fy - 1}). Turnover was {fmt(t, 'pct')} "
                     f"({fmt_delta(t - t_prev, 'pct')} vs FY{d.fy - 1}). {n_red} of {d.n_agencies} agencies have at "
                     f"least one red scorecard metric."),
    }
    red_t = [w for w in d.watch if w["metric"] == "turnover_rate"]
    if red_t:
        w = red_t[0]
        s = (f"{w['unit']} has the highest turnover at {fmt(w['value'], 'pct')}, {w['ratio']:.1f}x the overall "
             f"rate of {fmt(w['baseline'], 'pct')}.")
        if d.driver:
            hi, lo = d.driver["hi"], d.driver["lo"]
            s += (f" Within {d.driver['unit']}, turnover was {fmt(hi['value'], 'pct')} for {hi['pay_basis']} "
                  f"employees vs {fmt(lo['value'], 'pct')} for {lo['pay_basis']} employees.")
        out["concentration"] = s
    else:
        out["concentration"] = "No agency is at red status for turnover."
    if len(d.backtest):
        method, mape = forecast.select_method(d.backtest)
        mean = d.backtest.groupby("method")["mape"].mean()
        years = sorted(d.backtest.fiscal_year.unique())
        out["backtest"] = (f"Backtest FY{years[0]}-FY{years[-1]}, each year predicted from earlier years only, "
                           f"{d.n_agencies} agencies. Mean absolute percentage error: last-year rate "
                           f"{mean['last_rate']:.1%}, trailing 3-year rate {mean['trailing_mean']:.1%}, naive same "
                           f"count as last year {mean['naive_count']:.1%}. The projection uses the lowest "
                           f"({METHOD_LABEL[method]}); error of this size means it is a planning baseline, "
                           f"not a precise forecast.")
    out["definitions"] = " ".join(f"{m['label']}: {m['definition']}" for k, m in cat.items()
                                  if k in ("turnover_rate", "separations", "new_hire_attrition_rate"))
    return out


# ------------------------------------------------------------------ Excel

def write_excel(d: ReportData, path: Path) -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    cat, gov = settings.metric_catalog(), settings.governance()
    k = gov["privacy"]["min_cell_size"]
    number_formats = {"pct": "0.0%", "count": "#,##0", "usd": "$#,##0", "years": "0.0", "hours": "#,##0.0"}
    fills = {"red": "F6D5D5", "amber": "FDEFCB", "green": "D6EFD6"}
    bold, head_fill = Font(bold=True), PatternFill("solid", fgColor="EEEDE8")

    wb = Workbook()
    ws = wb.active
    ws.title = "About"
    about = [
        ("Workforce scorecard", f"FY{d.fy} (July 1, {d.fy - 1} - June 30, {d.fy})"),
        ("Generated", date.today().isoformat()),
        ("Data", "NYC Citywide Payroll Data (NYC Open Data, dataset k397-673e), " + (
            "SAMPLE MODE: a seeded random sample of real rows; counts describe the sample, not the workforce"
            if d.is_sample else f"{d.n_agencies} agencies in scope")),
        ("Benchmark", "BLS JOLTS total separations rate, state & local government excl. education (national)"),
        ("Privacy", f"Names are never downloaded. Cells with fewer than {k} employees show '{SUPPRESSED_TEXT.format(k=k)}'."),
        ("RAG status", f"Compared with the all-in-scope value: red >= {gov['scorecard']['red_ratio']}x, "
                       f"amber >= {gov['scorecard']['amber_ratio']}x (metrics where higher is worse)."),
        ("Definitions", "See the 'Metric definitions' sheet; all values come from config/metrics.yml."),
    ]
    for r in about:
        ws.append(r)
    for c in ws["A"]:
        c.font = bold
    ws.column_dimensions["A"].width, ws.column_dimensions["B"].width = 18, 120

    def card_sheet(title: str, card: scorecard.Scorecard):
        ws = wb.create_sheet(title)
        dim_label = settings.dimensions()[card.dimension]["label"]
        cols = [(card.dimension, dim_label, None)]
        for key in card.metric_keys:
            m = cat[key]
            cols.append((key, m["label"], m["format"]))
            cols.append((f"{key}__delta", "vs prior year", "pts" if m["format"] == "pct" else "chg"))
            if m["polarity"] != "neutral":
                cols.append((f"{key}__rag", "Status", "rag"))
        ws.append([c[1] for c in cols])
        for cell in ws[1]:
            cell.font, cell.fill = bold, head_fill
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        for _, rowdata in card.table.iterrows():
            ws.append([None] * len(cols))
            r = ws.max_row
            for j, (col, _, f) in enumerate(cols, start=1):
                cell = ws.cell(row=r, column=j)
                v = rowdata.get(col)
                if f is None:
                    cell.value = v
                elif f == "rag":
                    cell.value = charts.STATUS_LABEL.get(v or "", "")
                    if v in fills:
                        cell.fill = PatternFill("solid", fgColor=fills[v])
                elif pd.isna(v):
                    cell.value = SUPPRESSED_TEXT.format(k=k) if f in number_formats else None
                else:
                    cell.value = float(v)
                    cell.number_format = {"pts": '+0.0%;-0.0%', "chg": '+0.0%;-0.0%'}.get(f, number_formats.get(f))
            if rowdata[card.dimension] == scorecard.ALL_LABEL:
                for cell in ws[r]:
                    cell.font = bold
        ws.freeze_panes = "B2"
        ws.column_dimensions["A"].width = 36
        for j in range(2, len(cols) + 1):
            ws.column_dimensions[get_column_letter(j)].width = 14
        ws.row_dimensions[1].height = 45

    card_sheet("Agency scorecard", d.agency_card)
    card_sheet("Location scorecard", d.location_card)

    ws = wb.create_sheet("Trend")
    keys = list(cat)
    ws.append(["Fiscal year"] + [cat[k]["label"] for k in keys] + ["BLS benchmark: total separations rate"])
    bench = dict(zip(d.benchmark["fiscal_year"], d.benchmark["value"]))
    for _, r in d.trend.iterrows():
        ws.append([int(r["fiscal_year"])] + [None if pd.isna(r[k]) else float(r[k]) for k in keys]
                  + [bench.get(int(r["fiscal_year"]))])
        for j, key in enumerate(keys, start=2):
            ws.cell(row=ws.max_row, column=j).number_format = number_formats[cat[key]["format"]]
        ws.cell(row=ws.max_row, column=len(keys) + 2).number_format = "0.0%"
    for cell in ws[1]:
        cell.font, cell.fill = bold, head_fill
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[1].height = 45

    ws = wb.create_sheet("Watch list")
    ws.append(["Agency", "Metric", "Value", "All in scope", "Ratio"])
    for w in d.watch:
        f = cat[w["metric"]]["format"]
        ws.append([w["unit"], w["label"], float(w["value"]), float(w["baseline"]), round(float(w["ratio"]), 2)])
        ws.cell(row=ws.max_row, column=3).number_format = number_formats[f]
        ws.cell(row=ws.max_row, column=4).number_format = number_formats[f]
    for cell in ws[1]:
        cell.font, cell.fill = bold, head_fill

    ws = wb.create_sheet("Separations outlook")
    method = d.outlook["method"].iloc[0] if len(d.outlook) else ""
    ws.append([f"Projection for FY{d.fy + 1} using the backtest-selected method ({METHOD_LABEL.get(method, method)}). "
               "Planning estimate for whole agencies, not a prediction about individuals."])
    ws.append(["Agency", f"Headcount FY{d.fy}", f"Projected separations FY{d.fy + 1}", "Last-year rate method",
               "Trailing 3-year rate method", "Naive same count"])
    for _, r in d.outlook.iterrows():
        ws.append([r["agency"], float(r["headcount_prior"]), round(float(r["projected_separations"])),
                   round(float(r["last_rate"])), round(float(r["trailing_mean"])), round(float(r["naive_count"]))])
    ws.append([])
    ws.append(["Backtest (each year predicted from earlier years only)"])
    ws.append(["Fiscal year", "Method", "Agencies", "MAE (separations)", "MAPE", "Total predicted", "Total actual"])
    for _, r in d.backtest.iterrows():
        ws.append([int(r["fiscal_year"]), r["method"], int(r["agencies"]), round(r["mae"], 1), r["mape"],
                   round(r["total_predicted"]), round(r["total_actual"])])
        ws.cell(row=ws.max_row, column=5).number_format = "0.0%"
    ws.column_dimensions["A"].width = 36

    ws = wb.create_sheet("Metric definitions")
    ws.append(["Key", "Metric", "Definition", "Breakdowns", "Higher is", "Caveats"])
    for key, m in cat.items():
        ws.append([key, m["label"], m["definition"], ", ".join(m["dimensions"]),
                   {"higher_is_worse": "worse", "higher_is_better": "better"}.get(m["polarity"], "neutral"),
                   m.get("caveats", "")])
    for col, w in zip("ABCDEF", (24, 30, 70, 36, 10, 90)):
        ws.column_dimensions[col].width = w
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")

    ws = wb.create_sheet("Data quality")
    ws.append(["Status", "Check", "Category", "Failing rows", "Total rows", "Detail"])
    for _, r in d.dq.iterrows():
        ws.append([r["status"], r["check_name"], r["category"], int(r["failing_rows"]), int(r["total_rows"]), r["detail"]])
    ws.column_dimensions["B"].width, ws.column_dimensions["F"].width = 32, 140
    for sheet in wb.worksheets[4:]:
        for cell in sheet[1]:
            cell.font = bold

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


# ------------------------------------------------------------------ PowerPoint

def write_pptx(d: ReportData, path: Path) -> Path:
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.util import Inches, Pt

    ink, ink2 = RGBColor.from_string(charts.INK[1:]), RGBColor.from_string(charts.INK2[1:])
    cat = settings.metric_catalog()
    text = narrative(d)
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    blank = prs.slide_layouts[6]

    def add_text(slide, x, y, w, h, s, size=14, bold=False, color=ink2):
        tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = s
        p.font.size, p.font.bold, p.font.color.rgb = Pt(size), bold, color
        return tb

    def slide(title, subtitle=None):
        s = prs.slides.add_slide(blank)
        add_text(s, 0.6, 0.35, 12, 0.7, title, 26, True, ink)
        if subtitle:
            add_text(s, 0.6, 1.0, 12, 0.6, subtitle, 13)
        if d.is_sample:
            add_text(s, 10.3, 0.1, 3, 0.3, "SAMPLE DATA - not actual headcount", 10, True,
                     RGBColor.from_string(charts.STATUS["red"][1:]))
        return s

    def picture(s, png: bytes, x, y, w=None, h=None):
        s.shapes.add_picture(io.BytesIO(png), Inches(x), Inches(y), width=Inches(w) if w else None,
                             height=Inches(h) if h else None)

    # 1. title
    s = slide(f"Workforce scorecard, FY{d.fy}",
              f"{d.n_agencies} NYC agencies, public payroll records (NYC Open Data). Generated {date.today().isoformat()}.")
    add_text(s, 0.6, 2.2, 12, 2.0, text["headline"], 20, False, ink)

    # 2. KPI tiles
    s = slide("Headline indicators", f"FY{d.fy} vs FY{d.fy - 1}, all agencies in scope")
    tiles = ["headcount", "turnover_rate", "new_hire_attrition_rate", "ot_hours_per_worker", "median_base_salary"]
    for i, key in enumerate(tiles):
        cur, prev = d.headline[key]
        f = cat[key]["format"]
        delta = fmt_delta(cur - prev, "pct") if f == "pct" else f"{(cur / prev - 1) * 100:+.1f}%"
        x = 0.6 + i * 2.5
        add_text(s, x, 2.0, 2.3, 0.6, cat[key]["label"], 12)
        add_text(s, x, 2.5, 2.3, 0.8, fmt(cur, f), 28, True, ink)
        add_text(s, x, 3.3, 2.3, 0.5, f"{delta} vs FY{d.fy - 1}", 12)
    add_text(s, 0.6, 4.4, 12, 1.5, text["concentration"], 16, False, ink)

    # 3. trend vs benchmark
    s = slide("Turnover trend vs national benchmark",
              "In-scope turnover (separations / average headcount) and BLS JOLTS total separations rate for state & local "
              "government excl. education, summed over July-June. Context, not a like-for-like comparison.")
    tr = d.trend[["fiscal_year", "turnover_rate"]].rename(columns={"turnover_rate": "value"})
    picture(s, charts.png_trend({"In-scope agencies": tr, "BLS JOLTS (national)": d.benchmark}, "Annual turnover rate"),
            2.0, 1.75, h=5.4)

    # 4. concentration
    card = d.agency_card.table
    body = card[card["agency"] != scorecard.ALL_LABEL]
    s = slide("Where turnover is concentrated", text["concentration"])
    picture(s, charts.png_ranked_bars(body, "agency", "turnover_rate", f"Turnover rate by agency, FY{d.fy}", "pct",
                                      status_col="turnover_rate__rag", reference=d.agency_card.baseline["turnover_rate"]),
            2.4, 1.75, h=5.5)

    # 5. scorecard table
    s = slide("Agency scorecard", "Status compares each agency with the all-in-scope value "
              f"(red >= {settings.governance()['scorecard']['red_ratio']}x).")
    cols = [("agency", "Agency", None), ("headcount", "Headcount", "count"),
            ("turnover_rate", "Turnover", "pct"), ("turnover_rate__rag", "Status", "rag"),
            ("new_hire_attrition_rate", "New-hire attrition", "pct"), ("new_hire_attrition_rate__rag", "Status", "rag"),
            ("ot_hours_per_worker", "OT hrs / worker", "hours"), ("ot_hours_per_worker__rag", "Status", "rag")]
    rows = card.head(17)
    tbl = s.shapes.add_table(len(rows) + 1, len(cols), Inches(0.6), Inches(1.6), Inches(12.1), Inches(5.6)).table
    for j, (_, head, _) in enumerate(cols):
        tbl.cell(0, j).text = head
    k = settings.governance()["privacy"]["min_cell_size"]
    for i, (_, r) in enumerate(rows.iterrows(), start=1):
        for j, (col, _, f) in enumerate(cols):
            v = r.get(col)
            if f is None:
                val = str(v)
            elif f == "rag":
                val = charts.STATUS_LABEL.get(v or "", "")
            else:
                val = SUPPRESSED_TEXT.format(k=k) if pd.isna(v) else fmt(v, f).replace(" hours", "")
            tbl.cell(i, j).text = val
    for row in tbl.rows:
        row.height = Inches(0.3)
        for c in row.cells:
            for p in c.text_frame.paragraphs:
                p.font.size = Pt(10)
    tbl.columns[0].width = Inches(3.3)
    for j in range(1, len(cols)):
        tbl.columns[j].width = Inches((12.1 - 3.3) / (len(cols) - 1))

    # 6. location view
    loc = d.location_card.table
    loc_body = loc[loc["location"] != scorecard.ALL_LABEL]
    s = slide("Location view", "Turnover and overtime by work location (multi-site view).")
    picture(s, charts.png_ranked_bars(loc_body, "location", "turnover_rate", f"Turnover rate by location, FY{d.fy}", "pct",
                                      status_col="turnover_rate__rag", reference=d.location_card.baseline["turnover_rate"]),
            0.4, 1.8, 6.2)
    picture(s, charts.png_ranked_bars(loc_body, "location", "ot_hours_per_worker",
                                      f"Overtime hours per worker, FY{d.fy}", "hours",
                                      status_col="ot_hours_per_worker__rag",
                                      reference=d.location_card.baseline["ot_hours_per_worker"]), 6.8, 1.8, 6.2)

    # 7. outlook
    s = slide(f"Separations outlook, FY{d.fy + 1} (workforce planning)", text.get("backtest", ""))
    top = d.outlook.head(10)
    tbl = s.shapes.add_table(len(top) + 1, 3, Inches(0.6), Inches(2.0), Inches(8.5), Inches(4.2)).table
    for j, h in enumerate(["Agency", f"Headcount FY{d.fy}", f"Projected separations FY{d.fy + 1}"]):
        tbl.cell(0, j).text = h
    for i, (_, r) in enumerate(top.iterrows(), start=1):
        tbl.cell(i, 0).text = r["agency"]
        tbl.cell(i, 1).text = fmt(r["headcount_prior"], "count")
        tbl.cell(i, 2).text = fmt(r["projected_separations"], "count")
    method = d.outlook["method"].iloc[0] if len(d.outlook) else ""
    add_text(s, 9.4, 2.0, 3.5, 4, f"Method: {METHOD_LABEL.get(method, method)}, chosen because it had the lowest "
             "backtest error. A planning baseline for whole agencies, not a prediction about any individual.", 12)

    # 8. definitions & data notes
    s = slide("Definitions and data notes")
    notes = [
        text["definitions"],
        "Separations exclude CEASED records with zero regular hours (residual payments to earlier leavers).",
        "Source does not distinguish voluntary from involuntary exits, and has no engagement, performance, or "
        "demographic data.",
        f"Privacy: names are never downloaded; any cell with fewer than {k} employees is suppressed.",
        "Data: NYC Citywide Payroll Data (NYC Open Data); BLS JOLTS via the BLS public API.",
    ]
    add_text(s, 0.6, 1.4, 12, 5.5, "\n\n".join(notes), 13)

    path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(path)
    return path


def build_all(con, fy: int, out_dir: Path, is_sample: bool = False) -> tuple[Path, Path]:
    d = collect(con, fy, is_sample)
    suffix = "_SAMPLE" if is_sample else ""
    return (write_excel(d, out_dir / f"workforce_scorecard_FY{fy}{suffix}.xlsx"),
            write_pptx(d, out_dir / f"workforce_scorecard_FY{fy}{suffix}.pptx"))
