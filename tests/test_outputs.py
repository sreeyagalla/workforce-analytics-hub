"""Scorecard, planning backtest, Excel/PowerPoint exports, and the dashboard itself."""
from __future__ import annotations

from pathlib import Path

import numpy as np
import pandas as pd
import pytest

import reference
from wfa import forecast, reports, scorecard, settings
from wfa.metrics import fmt

ROOT = Path(__file__).resolve().parents[1]
FY = 2025


def test_rag_status_rules():
    r = scorecard.rag_status
    assert r(1.30, 1.0, "higher_is_worse", 1.10, 1.25) == "red"
    assert r(1.15, 1.0, "higher_is_worse", 1.10, 1.25) == "amber"
    assert r(1.00, 1.0, "higher_is_worse", 1.10, 1.25) == "green"
    assert r(5.0, 1.0, "neutral", 1.10, 1.25) == ""
    assert r(np.nan, 1.0, "higher_is_worse", 1.10, 1.25) == ""


def test_agency_scorecard_matches_independent_computation(sample_con, sample_ref):
    card = scorecard.build(sample_con, FY, "agency")
    gov = settings.governance()["scorecard"]
    for key in ["headcount", "turnover_rate", "new_hire_attrition_rate", "ot_hours_per_worker"]:
        ref = reference.compute(sample_ref, key, "agency", None, [FY]).set_index("agency")["value"]
        total = reference.compute(sample_ref, key, None, None, [FY])["value"].iloc[0]
        got = card.table.set_index("agency")[key]
        assert got[scorecard.ALL_LABEL] == pytest.approx(total)
        for agency, v in ref.items():
            assert got[agency] == pytest.approx(v, nan_ok=True)
            if settings.metric_catalog()[key]["polarity"] == "higher_is_worse":
                expected = scorecard.rag_status(v, total, "higher_is_worse", gov["amber_ratio"], gov["red_ratio"])
                assert card.table.set_index("agency").loc[agency, f"{key}__rag"] == expected


def test_backtest_uses_only_prior_year_inputs(sample_con, sample_ref):
    preds, summary = forecast.backtest(sample_con)
    t = reference.compute(sample_ref, "turnover_rate", "agency").set_index(["fiscal_year", "agency"])
    h = reference.compute(sample_ref, "headcount", "agency").set_index(["fiscal_year", "agency"])
    for _, p in preds.iterrows():
        prior = (p.fiscal_year - 1, p.agency)
        assert p.last_rate == pytest.approx(h.loc[prior, "value"] * t.loc[prior, "value"])
        assert p.naive_count == t.loc[prior, "numerator"]
        assert p.actual == t.loc[(p.fiscal_year, p.agency), "numerator"]
    assert set(summary["method"]) == set(forecast.METHODS)


def test_outlook_uses_backtest_selected_method(sample_con):
    _, summary = forecast.backtest(sample_con)
    method, _ = forecast.select_method(summary)
    out = forecast.outlook(sample_con)
    assert (out["method"] == method).all()
    assert (out["projected_separations"] == out[method]).all()


@pytest.fixture(scope="module")
def report_data(sample_con):
    return reports.collect(sample_con, FY, is_sample=True)


def test_excel_scorecard_values_match_independent_computation(report_data, sample_ref, tmp_path):
    from openpyxl import load_workbook
    path = reports.write_excel(report_data, tmp_path / "s.xlsx")
    wb = load_workbook(path)
    assert wb.sheetnames == ["About", "Agency scorecard", "Location scorecard", "Trend", "Watch list",
                             "Separations outlook", "Metric definitions", "Data quality"]
    ws = wb["Agency scorecard"]
    header = [c.value for c in ws[1]]
    col = header.index("Turnover rate")
    ref = reference.compute(sample_ref, "turnover_rate", "agency", None, [FY]).set_index("agency")["value"]
    checked = 0
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[0] in ref.index:
            assert row[col] == pytest.approx(ref[row[0]])
            checked += 1
    assert checked == len(ref)
    assert "SAMPLE MODE" in wb["About"]["B3"].value


def test_excel_shows_suppressed_cells_as_text(sample_con, tmp_path):
    from openpyxl import load_workbook
    d = reports.collect(sample_con, FY, is_sample=True)
    path = reports.write_excel(d, tmp_path / "s.xlsx")
    ws = load_workbook(path)["Location scorecard"]
    values = [c for row in ws.iter_rows(min_row=2, values_only=True) for c in row]
    k = settings.governance()["privacy"]["min_cell_size"]
    card = d.location_card.table
    if card[d.location_card.metric_keys].isna().any().any():
        assert f"n<{k}" in values


def test_pptx_deck_is_built_with_grounded_headline(report_data, sample_ref, tmp_path):
    from pptx import Presentation
    path = reports.write_pptx(report_data, tmp_path / "s.pptx")
    prs = Presentation(path)
    assert len(prs.slides) == 8
    text = " ".join(sh.text_frame.text for s in prs.slides for sh in s.shapes if sh.has_text_frame)
    hc = reference.compute(sample_ref, "headcount", None, None, [FY])["value"].iloc[0]
    to = reference.compute(sample_ref, "turnover_rate", None, None, [FY])["value"].iloc[0]
    assert f"End-of-year headcount was {fmt(hc, 'count')} in FY{FY}" in text
    assert f"Turnover was {fmt(to, 'pct')}" in text
    assert "SAMPLE DATA" in text


def test_dashboard_renders_every_tab_and_answers_a_question(sample_paths, monkeypatch):
    from streamlit.testing.v1 import AppTest
    monkeypatch.setenv("WFA_SAMPLE", "1")
    monkeypatch.setenv("WFA_WAREHOUSE", str(sample_paths.warehouse))
    monkeypatch.delenv("ANTHROPIC_API_KEY", raising=False)
    at = AppTest.from_file(str(ROOT / "app" / "streamlit_app.py"), default_timeout=180)
    at.run()
    assert not at.exception, at.exception
    assert at.title[0].value == "Workforce Analytics Hub"
    assert [t.label for t in at.tabs] == ["Overview", "Scorecards", "Explore", "Ask a question", "Planning",
                                          "Data quality", "Definitions"]
    assert len(at.metric) >= 5
    at.text_input(key="question").input("What was the turnover rate at Parks in FY2025?").run()
    assert not at.exception
    assert any("Turnover rate for Parks & Recreation in FY2025" in s.value for s in at.success)
    at.text_input(key="question").input("What is John Smith's salary?").run()
    assert any("only answer aggregate questions" in w.value for w in at.warning)
